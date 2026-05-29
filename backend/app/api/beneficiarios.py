"""API do cadastro mestre de beneficiários (prestadores).

Operações:
    - Listar (com filtros) / detalhar
    - Criar manual / atualizar / aprovar / desativar
    - Importar planilha em 2 passos (preview → confirma)

Permissão:
    - Operações de leitura: qualquer usuário com visão executiva
      (admin, aprovador, operador). Coordenador NÃO acessa.
    - Operações de escrita: ADMIN ou APROVADOR (no fluxo do MedPag,
      cadastrar prestador é decisão de negócio que requer responsabilidade
      pela qualidade dos dados bancários).
"""

from __future__ import annotations

import io
from uuid import UUID

import structlog
from fastapi import (
    APIRouter,
    Depends,
    File,
    Form,
    HTTPException,
    Query,
    Response,
    UploadFile,
    status,
)
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.deps import (
    get_db,
    require_aprovador,
    require_visao_executiva,
)
from app.core.exceptions import ValidacaoError
from app.models.beneficiario import StatusBeneficiario
from app.models.user import User
from app.schemas.beneficiario import (
    BeneficiarioCreateRequest,
    BeneficiarioListResponse,
    BeneficiarioOut,
    BeneficiarioUpdateRequest,
    ImportConfirmRequest,
    ImportConfirmResponse,
    ImportPreviewResponse,
)
from app.services.beneficiario_service import (
    BeneficiarioNaoEncontradoError,
    BeneficiarioService,
    BeneficiarioServiceError,
    ClienteNaoEncontradoError,
)

log = structlog.get_logger()

router = APIRouter(prefix="/api/beneficiarios", tags=["beneficiarios"])


# ============================================================
# Listagem e detalhe
# ============================================================


@router.get("", response_model=BeneficiarioListResponse)
async def listar_beneficiarios(
    cliente_id: UUID | None = Query(default=None),
    status_filter: StatusBeneficiario | None = Query(default=None, alias="status"),
    search: str | None = Query(default=None, max_length=120),
    page: int = Query(default=1, ge=1),
    per_page: int = Query(default=50, ge=1, le=200),
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_visao_executiva),
) -> BeneficiarioListResponse:
    # Multi-tenancy: força filtro do cliente do user
    if user.cliente_id is not None:
        cliente_id = user.cliente_id
    service = BeneficiarioService(db)
    return await service.listar(
        cliente_id=cliente_id,
        status=status_filter,
        search=search,
        page=page,
        per_page=per_page,
    )


@router.get("/template.xlsx")
async def baixar_template_planilha(
    _: User = Depends(require_visao_executiva),
) -> StreamingResponse:
    """Baixa um template XLSX com os cabecalhos esperados e 2 linhas de exemplo.

    Util pra dar ao hospital uma planilha pronta de preencher antes de
    importar a base. Os cabecalhos batem com os aliases reconhecidos
    em `BeneficiarioService._ALIASES_BENEFICIARIO` (case + acento
    insensitive).
    """
    wb = Workbook()
    ws = wb.active
    if ws is None:
        ws = wb.create_sheet("Prestadores")
    else:
        ws.title = "Prestadores"

    cabecalhos = [
        "CPF",
        "Nome",
        "CRM",
        "Categoria",
        "Especialidade",
        "Email",
        "Telefone",
        "Banco",
        "Agencia",
        "Conta",
        "Tipo PIX",
        "Chave PIX",
        "Valor padrao (R$)",
        "Observacoes",
    ]

    fonte_cab = Font(bold=True, color="FFFFFF")
    preench_cab = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    centralizar = Alignment(horizontal="center")

    for col_idx, cab in enumerate(cabecalhos, start=1):
        cell = ws.cell(row=1, column=col_idx, value=cab)
        cell.font = fonte_cab
        cell.fill = preench_cab
        cell.alignment = centralizar

    # Larguras razoaveis pra leitura
    larguras = [16, 32, 12, 18, 22, 30, 16, 8, 10, 14, 12, 32, 16, 32]
    for i, w in enumerate(larguras, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

    # Linhas de exemplo (vamos deixar 2 pra ficar claro o formato)
    exemplos = [
        [
            "123.456.789-09",
            "MARIA SOUZA DA SILVA",
            "CRM-RJ 123456",
            "Medico",
            "Cardiologia",
            "maria.silva@exemplo.com.br",
            "(21) 98765-4321",
            "341",
            "1234",
            "56789-0",
            "CPF",
            "12345678909",
            "1500.00",
            "Plantao 12h fixo",
        ],
        [
            "987.654.321-00",
            "JOAO PEREIRA SANTOS",
            "CRM-RJ 654321",
            "Medico",
            "Cirurgia Geral",
            "joao.santos@exemplo.com.br",
            "(21) 99999-0000",
            "237",
            "0001",
            "12345-6",
            "EMAIL",
            "joao.santos@exemplo.com.br",
            "1800.00",
            "",
        ],
    ]
    for r_idx, linha in enumerate(exemplos, start=2):
        for c_idx, valor in enumerate(linha, start=1):
            ws.cell(row=r_idx, column=c_idx, value=valor)

    # Aba auxiliar com instrucoes (ajuda muito quando alguem nao tecnico abre)
    ws_help = wb.create_sheet("Instrucoes")
    ws_help.column_dimensions["A"].width = 22
    ws_help.column_dimensions["B"].width = 90

    instrucoes = [
        ("Coluna", "Como preencher"),
        ("CPF", "Obrigatorio. Aceita 000.000.000-00 ou 00000000000."),
        ("Nome", "Obrigatorio. Nome completo do prestador."),
        ("CRM", "Opcional. Conselho profissional (CRM, COREN, CRO etc)."),
        ("Categoria", "Opcional. Ex: Medico, Enfermeiro, Limpeza, RH."),
        ("Especialidade", "Opcional. Ex: Cardiologia, UTI, Pediatria."),
        ("Email", "Opcional."),
        ("Telefone", "Opcional. (DDD) numero, com ou sem mascara."),
        (
            "Banco",
            "Codigo FEBRABAN de 3 digitos. Ex: 341 (Itau), 237 (Bradesco), 136 (Unicred), 260 (Nubank), 077 (Inter). Codigo invalido = linha rejeitada.",
        ),
        ("Agencia", "Numero da agencia (4-5 digitos)."),
        ("Conta", "Numero da conta com digito (ex: 12345-6)."),
        (
            "Tipo PIX",
            "CPF, CNPJ, EMAIL, TELEFONE ou ALEATORIA. Se tipo=CPF, a chave precisa ser o MESMO CPF do prestador (anti-fraude).",
        ),
        ("Chave PIX", "A chave em si. Para CPF/TELEFONE pode ser sem mascara."),
        (
            "Valor padrao (R$)",
            "Opcional. Aceita 1500.00 ou 1500,00 ou 150000 (centavos).",
        ),
        ("Observacoes", "Texto livre."),
        ("", ""),
        ("DICA", "Se o CPF ja existir no cadastro, a planilha ATUALIZA o registro."),
        (
            "DICA",
            "Caso voce envie a planilha sem dados bancarios, o prestador entra como ATIVO mas nao recebe pagamento ate completar.",
        ),
        (
            "DICA",
            "Pagamento por PIX e por TED sao independentes - pode preencher SO PIX, SO banco, ou os DOIS (PIX e' priorizado).",
        ),
        (
            "VALIDACAO",
            "A plataforma valida: CPF com digito, Banco contra a lista oficial FEBRABAN, e PIX-CPF tem que bater com o CPF do prestador.",
        ),
        (
            "VALIDACAO",
            "Linhas com erro aparecem em vermelho na tela de pre-visualizacao - voce pode corrigir e re-enviar.",
        ),
    ]
    for r_idx, (a, b) in enumerate(instrucoes, start=1):
        cell_a = ws_help.cell(row=r_idx, column=1, value=a)
        cell_b = ws_help.cell(row=r_idx, column=2, value=b)
        if r_idx == 1:
            cell_a.font = Font(bold=True)
            cell_b.font = Font(bold=True)
        if a == "DICA":
            cell_a.font = Font(bold=True, color="B45309")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    headers = {
        "Content-Disposition": 'attachment; filename="medpag_prestadores_template.xlsx"',
    }
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers,
    )


@router.get("/{beneficiario_id}", response_model=BeneficiarioOut)
async def detalhar_beneficiario(
    beneficiario_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_visao_executiva),
) -> BeneficiarioOut:
    from app.core.deps import verificar_acesso_cliente

    service = BeneficiarioService(db)
    try:
        b = await service.buscar(beneficiario_id)
    except BeneficiarioNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    verificar_acesso_cliente(
        user,
        b.cliente_id,
        mensagem="Beneficiário pertence a outro cliente.",
    )
    return BeneficiarioOut.model_validate(b)


# ============================================================
# Criar / atualizar / aprovar / desativar
# ============================================================


@router.post(
    "",
    response_model=BeneficiarioOut,
    status_code=status.HTTP_201_CREATED,
)
async def criar_beneficiario(
    payload: BeneficiarioCreateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_aprovador),
) -> BeneficiarioOut:
    service = BeneficiarioService(db)
    try:
        b = await service.criar(payload)
    except ClienteNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except BeneficiarioServiceError as e:
        raise ValidacaoError(str(e)) from e
    log.info(
        "beneficiario.criado",
        beneficiario_id=str(b.id),
        cliente_id=str(b.cliente_id),
        criador=user.email,
    )
    return BeneficiarioOut.model_validate(b)


@router.patch("/{beneficiario_id}", response_model=BeneficiarioOut)
async def atualizar_beneficiario(
    beneficiario_id: UUID,
    payload: BeneficiarioUpdateRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_aprovador),
) -> BeneficiarioOut:
    service = BeneficiarioService(db)
    try:
        b = await service.atualizar(beneficiario_id, payload)
    except BeneficiarioNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except BeneficiarioServiceError as e:
        raise ValidacaoError(str(e)) from e
    log.info(
        "beneficiario.atualizado",
        beneficiario_id=str(b.id),
        atualizador=user.email,
    )
    return BeneficiarioOut.model_validate(b)


@router.post("/{beneficiario_id}/aprovar", response_model=BeneficiarioOut)
async def aprovar_beneficiario(
    beneficiario_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_aprovador),
) -> BeneficiarioOut:
    service = BeneficiarioService(db)
    try:
        b = await service.aprovar(beneficiario_id)
    except BeneficiarioNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    log.info(
        "beneficiario.aprovado",
        beneficiario_id=str(b.id),
        aprovador=user.email,
    )
    return BeneficiarioOut.model_validate(b)


@router.post("/{beneficiario_id}/desativar", response_model=BeneficiarioOut)
async def desativar_beneficiario(
    beneficiario_id: UUID,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_aprovador),
) -> BeneficiarioOut:
    service = BeneficiarioService(db)
    try:
        b = await service.desativar(beneficiario_id)
    except BeneficiarioNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    log.info(
        "beneficiario.desativado",
        beneficiario_id=str(b.id),
        ator=user.email,
    )
    return BeneficiarioOut.model_validate(b)


# ============================================================
# Importação em massa (XLSX/CSV)
# ============================================================


@router.post(
    "/import/preview",
    response_model=ImportPreviewResponse,
)
async def preview_importacao(
    cliente_id: UUID = Form(...),
    arquivo: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    _: User = Depends(require_aprovador),
) -> ImportPreviewResponse:
    """Recebe a planilha e devolve um relatório linha-a-linha com erros e
    avisos. Nada é gravado ainda — usuário revisa e chama /import/confirm.
    """
    if not arquivo.filename:
        raise ValidacaoError("Nome do arquivo é obrigatório.")
    conteudo = await arquivo.read()
    if not conteudo:
        raise ValidacaoError("Arquivo vazio.")

    service = BeneficiarioService(db)
    try:
        return await service.importar_preview(
            cliente_id=cliente_id,
            conteudo=conteudo,
            nome_arquivo=arquivo.filename,
        )
    except ClienteNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except BeneficiarioServiceError as e:
        raise ValidacaoError(str(e)) from e


@router.post(
    "/import/confirm",
    response_model=ImportConfirmResponse,
)
async def confirmar_importacao(
    payload: ImportConfirmRequest,
    db: AsyncSession = Depends(get_db),
    user: User = Depends(require_aprovador),
) -> ImportConfirmResponse:
    """Confirma e efetiva a importação previamente analisada via /preview."""
    service = BeneficiarioService(db)
    try:
        resultado = await service.importar_confirmar(payload)
    except ClienteNaoEncontradoError as e:
        raise HTTPException(status_code=404, detail=str(e)) from e
    except BeneficiarioServiceError as e:
        raise ValidacaoError(str(e)) from e

    log.info(
        "beneficiario.import.confirmado",
        criados=resultado.qtd_criados,
        atualizados=resultado.qtd_atualizados,
        ignorados=resultado.qtd_ignorados,
        erros=resultado.qtd_erros,
        ator=user.email,
    )
    return resultado


__all__ = ["router"]
