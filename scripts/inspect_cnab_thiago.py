"""Inspeciona o template CNAB Unicred enviado pelo Thiago.

Output esperado: estrutura das abas relevantes pra entender
quais campos a planilha aceita por modalidade (TED/PIX/etc).
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl

CAMINHO = Path("data/thiago/cnab-mestre-auris.xlsm")

ABAS_INTERESSE = [
    "INICIO",
    "CNAB",
    "AUX",
]


def main() -> None:
    wb = openpyxl.load_workbook(CAMINHO, data_only=True, keep_vba=False)

    for aba in ABAS_INTERESSE:
        if aba not in wb.sheetnames:
            continue
        ws = wb[aba]
        print(f"=========== {aba}  ({ws.max_row} linhas x {ws.max_column} colunas) ===========")
        max_show = ws.max_row
        for i, row in enumerate(
            ws.iter_rows(min_row=1, max_row=max_show, values_only=True), start=1
        ):
            cells = []
            for c in row:
                if c is None:
                    cells.append("")
                else:
                    cells.append(str(c)[:60])
            print(f"  L{i:2d}: " + " | ".join(cells))
        print()


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"ERRO: {exc}", file=sys.stderr)
        sys.exit(1)
